import docx
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Margens padrão A4 / Carta
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

COLOR_PRIMARY = RGBColor(31, 78, 121)    # Navy Blue
COLOR_SECONDARY = RGBColor(89, 89, 89)   # Slate Gray
COLOR_TEXT = RGBColor(38, 38, 38)        # Charcoal
COLOR_ENUNCIADO = RGBColor(46, 117, 182) # Accent Blue

def set_run_font(run, name="Calibri", size_pt=11, color=COLOR_TEXT, bold=False, italic=False):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic

def add_header_title(doc, title, subtitle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(title)
    set_run_font(run_t, name="Calibri", size_pt=18, color=COLOR_PRIMARY, bold=True)
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(14)
    run_s = p2.add_run(subtitle)
    set_run_font(run_s, name="Calibri", size_pt=11, color=COLOR_SECONDARY, italic=True)
    
    p_hr = doc.add_paragraph()
    p_hr.paragraph_format.space_after = Pt(14)
    p_hr_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                            r'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="1F4E79"/>'
                            r'</w:pBdr>')
    p_hr._p.get_or_add_pPr().append(p_hr_border)

def add_code_block(doc, code_str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    
    shading = parse_xml(r'<w:shd {} w:fill="F4F5F7"/>'.format(nsdecls('w')))
    cell._tc.get_or_add_tcPr().append(shading)
    
    borders = parse_xml(r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
                        r'<w:left w:val="single" w:sz="18" w:space="0" w:color="1F4E79"/>'
                        r'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
                        r'<w:right w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>'
                        r'</w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(borders)
    
    cell.paragraphs[0].text = ""
    lines = code_str.strip().split("\n")
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size_pt=9.5, color=RGBColor(33, 37, 41))
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(6)

def add_question(doc, num, enunciado, resposta_paragraphs, code_snippets=None, table_data=None, note=None):
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(14)
    p_q.paragraph_format.space_after = Pt(3)
    p_q.paragraph_format.keep_with_next = True
    run_q = p_q.add_run(f"Questão {num}")
    set_run_font(run_q, name="Calibri", size_pt=13, color=COLOR_PRIMARY, bold=True)
    
    p_e = doc.add_paragraph()
    p_e.paragraph_format.space_after = Pt(6)
    p_e.paragraph_format.left_indent = Inches(0.2)
    p_e.paragraph_format.keep_with_next = True
    run_el = p_e.add_run("Enunciado: ")
    set_run_font(run_el, name="Calibri", size_pt=11, color=COLOR_ENUNCIADO, bold=True)
    run_et = p_e.add_run(enunciado)
    set_run_font(run_et, name="Calibri", size_pt=11, color=COLOR_ENUNCIADO, italic=True)
    
    p_r = doc.add_paragraph()
    p_r.paragraph_format.space_after = Pt(4)
    p_r.paragraph_format.keep_with_next = True
    run_rl = p_r.add_run("Resposta:")
    set_run_font(run_rl, name="Calibri", size_pt=11, color=COLOR_PRIMARY, bold=True)
    
    for item in resposta_paragraphs:
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_after = Pt(5)
        p_b.paragraph_format.line_spacing = 1.15
        if isinstance(item, tuple) and len(item) == 2:
            prefix, text = item
            run_p = p_b.add_run(prefix)
            set_run_font(run_p, name="Calibri", size_pt=11, color=COLOR_TEXT, bold=True)
            run_t = p_b.add_run(text)
            set_run_font(run_t, name="Calibri", size_pt=11, color=COLOR_TEXT)
        else:
            run_t = p_b.add_run(item)
            set_run_font(run_t, name="Calibri", size_pt=11, color=COLOR_TEXT)
            
    if table_data:
        headers, rows = table_data
        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for col_idx, h in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = h
            shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
            cell._tc.get_or_add_tcPr().append(shading)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, name="Calibri", size_pt=10.5, color=RGBColor(255, 255, 255), bold=True)
                    
        for row_idx, row in enumerate(rows):
            for col_idx, val in enumerate(row):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = val
                bg = "F9FAFB" if row_idx % 2 == 0 else "FFFFFF"
                shading = parse_xml(r'<w:shd {} w:fill="{}"/>'.format(nsdecls('w'), bg))
                cell._tc.get_or_add_tcPr().append(shading)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    for run in p.runs:
                        set_run_font(run, name="Calibri", size_pt=10, color=COLOR_TEXT)
        
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(6)

    if code_snippets:
        for code in code_snippets:
            add_code_block(doc, code)

    if note:
        p_n = doc.add_paragraph()
        p_n.paragraph_format.space_before = Pt(3)
        p_n.paragraph_format.space_after = Pt(8)
        p_n.paragraph_format.left_indent = Inches(0.2)
        run_nl = p_n.add_run("Referência / Código-Fonte: ")
        set_run_font(run_nl, name="Calibri", size_pt=10, color=COLOR_SECONDARY, bold=True)
        run_nt = p_n.add_run(note)
        set_run_font(run_nt, name="Calibri", size_pt=10, color=COLOR_SECONDARY, italic=True)

# CABEÇALHO DO DOCUMENTO
add_header_title(
    doc,
    "Trabalho Prático 1 (TP1) - Arquitetura de Software e DDD",
    "Engenharia Disciplinada de Software | Perguntas e Respostas Objetivas"
)

# QUESTÃO 1
add_question(
    doc, 1,
    "Explique de forma sucinta qual a razão de criar Agregados.",
    [
        "No Domain-Driven Design (DDD), a principal razão para criar Agregados (Aggregates) é delimitar fronteiras claras de consistência transacional e encapsular as invariantes de negócio.",
        "Em vez de permitir que o modelo de domínio seja tratado como uma teia desordenada de entidades interligadas em memória, o Agregado agrupa entidades e objetos de valor que devem mudar juntos. O acesso externo ocorre exclusivamente através de uma entidade líder, a Raiz do Agregado (Aggregate Root). Com isso, garante-se que nenhuma alteração interna ocorra sem que as regras do domínio sejam integralmente validadas em uma única transação atômica, evitando locks concorrentes e corrupção de dados."
    ]
)

# QUESTÃO 2
add_question(
    doc, 2,
    "O que significa “consistência transacional”?",
    [
        "Consistência transacional é a garantia de que uma transação (conjunto de operações de leitura e escrita) levará os dados do sistema de um estado válido inicial para outro estado igualmente válido, respeitando de forma imediata e estrita todas as regras de integridade e invariantes de negócio.",
        "Caso qualquer operação do bloco transacional viole uma regra de negócio ou sofra falha técnica, ocorre o rollback imediato, impedindo que dados intermediários, parciais ou inválidos persistam no banco de dados. No contexto do DDD, o limite de um Agregado define exatamente o limite da consistência transacional: cada transação deve modificar apenas uma instância de Agregado."
    ]
)

# QUESTÃO 3
add_question(
    doc, 3,
    "Cite as 4 propriedades cruciais que definem transações.",
    [
        "As 4 propriedades essenciais que definem transações compõem o acrônimo ACID:",
        ("1. Atomicidade (Atomicity): ", "Princípio do 'tudo ou nada'. Todas as operações contidas na transação são executadas com sucesso e confirmadas, ou nenhuma alteração é aplicada, revertendo qualquer modificação intermediária em caso de falha."),
        ("2. Consistência (Consistency): ", "Garante que a transação preserve todas as invariantes e regras de validação do domínio, transitando exclusivamente entre estados válidos e em conformidade com o modelo."),
        ("3. Isolamento (Isolation): ", "Garante que transações executadas concorrentemente não interfiram umas nas outras. O resultado visível no banco de dados deve ser idêntico ao que seria obtido se fossem executadas de maneira estritamente sequencial."),
        ("4. Durabilidade (Durability): ", "Assegura que, uma vez confirmado (commit) o término da transação, as alterações tornam-se permanentes e não serão perdidas, resistindo a falhas de hardware, falta de energia ou reinicializações do servidor.")
    ]
)

# QUESTÃO 4
add_question(
    doc, 4,
    "O que são “invariantes de negócio”?",
    [
        "Invariantes de negócio são regras, condições e restrições lógicas do domínio que devem permanecer verdadeiras em todos os momentos durante o ciclo de vida dos dados e das operações do sistema.",
        "Elas representam a coerência fundamental do modelo: se uma invariante for violada, o estado do sistema é considerado inconsistente e corrompido.",
        ("Exemplos práticos: ", "Um pedido não pode ser finalizado/pago sem ao menos um item associado; a quantidade de um item no pedido deve ser sempre maior que zero; o saldo de uma conta corrente sem limite não pode ser negativo."),
        "No DDD, a Raiz do Agregado é a única responsável por verificar e garantir que nenhuma invariante seja violada antes de persistir qualquer transição de estado."
    ]
)

# QUESTÃO 5
add_question(
    doc, 5,
    "Porque um agregado só deve ter acesso a outro agregado pelo ID?",
    [
        "Um agregado só deve referenciar outro agregado através do seu identificador único (ID), e não por referência direta de objeto em memória (como @OneToOne ou instâncias injetadas diretamente), pelas seguintes razões arquiteturais fundamentais:",
        ("1. Preservação dos limites transacionais: ", "A regra de ouro do DDD estabelece que uma única transação deve modificar apenas uma instância de Agregado. Referências diretas em memória estimulam desenvolvedores a carregar e modificar múltiplos agregados no mesmo commit, quebrando o isolamento."),
        ("2. Escalabilidade e persistência desacoplada: ", "Ao armazenar apenas o ID, os agregados podem residir facilmente em tabelas distintas, bancos de dados diferentes ou microsserviços autônomos, sem acoplamento de schema relacional."),
        ("3. Desempenho e gestão de memória: ", "Evita o problema de carregar grafos gigantescos de objetos interdependentes em cascata (lazy/eager loading descontrolado), otimizando consultas SQL e economizando memória."),
        ("4. Respeito estrito ao encapsulamento: ", "Impede que um agregado acesse métodos internos de outro agregado sem passar pelo contrato formal de fronteira.")
    ]
)

# QUESTÃO 6
code_q6 = """package br.edu.infnet.domain.model;

import java.math.BigDecimal;
import java.util.ArrayList;
import java.util.Collections;
import java.util.List;
import java.util.Objects;
import java.util.UUID;

public class Pedido {

    private final UUID id;
    
    // Referência a outro agregado (Cliente) estritamente pelo ID:
    private final UUID clienteId;
    
    private final List<ItemPedido> itens = new ArrayList<>();
    private StatusPedido status;
    private BigDecimal valorTotal;

    public Pedido(UUID id, UUID clienteId) {
        this.id = Objects.requireNonNull(id, "ID do pedido é obrigatório.");
        this.clienteId = Objects.requireNonNull(clienteId, "ClienteID é obrigatório.");
        this.status = StatusPedido.CRIADO;
        this.valorTotal = BigDecimal.ZERO;
    }

    public void adicionarItem(UUID produtoId, String descricao, int quantidade, BigDecimal precoUnitario) {
        if (this.status != StatusPedido.CRIADO) {
            throw new IllegalStateException("Não é permitido adicionar itens em pedidos finalizados.");
        }
        ItemPedido item = new ItemPedido(produtoId, descricao, quantidade, precoUnitario);
        this.itens.add(item);
        this.valorTotal = this.valorTotal.add(item.calcularSubtotal());
    }

    public UUID getId() { return id; }
    public UUID getClienteId() { return clienteId; }
    public StatusPedido getStatus() { return status; }
    public BigDecimal getValorTotal() { return valorTotal; }
    public List<ItemPedido> getItens() { return Collections.unmodifiableList(itens); }
}"""

add_question(
    doc, 6,
    "Crie um trecho de código Java de uma entidade que represente um agregado e faça referência a outro agregado dentro do escopo do projeto Pet Friends.",
    [
        "Abaixo está a implementação da classe Pedido, que atua como Raiz do Agregado (Aggregate Root). Ela faz referência ao agregado Cliente exclusivamente através do identificador clienteId, mantendo seus itens e regras internas devidamente encapsulados:"
    ],
    code_snippets=[code_q6],
    note="https://github.com/Marcus-Boni/INFNET-26E2-26E3/blob/main/26E3_4/TP1/src/main/java/br/edu/infnet/domain/model/Pedido.java"
)

# QUESTÃO 7
code_q7 = """    // Lista interna de eventos de domínio pendentes de publicação
    private final List<DomainEvent> domainEvents = new ArrayList<>();

    /**
     * Método de negócio com validação de invariantes e previsão de publicação de evento de domínio.
     */
    public void pagar(BigDecimal valorPago) {
        // 1. Validação de Invariantes de Negócio
        if (this.status == StatusPedido.PAGO) {
            throw new IllegalStateException("O pedido já se encontra pago.");
        }
        if (this.status == StatusPedido.CANCELADO) {
            throw new IllegalStateException("Não é possível pagar um pedido previamente cancelado.");
        }
        if (this.itens.isEmpty()) {
            throw new IllegalStateException("O pedido não pode ser pago pois não possui itens.");
        }
        if (valorPago == null || valorPago.compareTo(this.valorTotal) < 0) {
            throw new IllegalArgumentException("O valor pago é insuficiente para quitar o total do pedido.");
        }

        // 2. Mudança de estado interno do agregado
        this.status = StatusPedido.PAGO;

        // 3. Previsão e registro do evento de domínio para futura publicação
        PedidoPagoEvent evento = new PedidoPagoEvent(this.id, this.clienteId, valorPago);
        this.domainEvents.add(evento);
    }

    public List<DomainEvent> pullDomainEvents() {
        List<DomainEvent> recordedEvents = new ArrayList<>(this.domainEvents);
        this.domainEvents.clear();
        return Collections.unmodifiableList(recordedEvents);
    }"""

add_question(
    doc, 7,
    "Elabore um trecho de código Java que mostre um método de negócio com a previsão de publicação de um evento de domínio dentro do escopo do projeto Pet Friends.",
    [
        "No agregado Pedido, o método de negócio pagar valida as invariantes de negócio, altera o estado interno e registra a ocorrência do evento de domínio PedidoPagoEvent na lista interna de eventos pendentes de publicação:"
    ],
    code_snippets=[code_q7],
    note="https://github.com/Marcus-Boni/INFNET-26E2-26E3/blob/main/26E3_4/TP1/src/main/java/br/edu/infnet/domain/model/Pedido.java (linhas 45 a 70)"
)

# QUESTÃO 8
add_question(
    doc, 8,
    "O que é “evento de domínio”?",
    [
        "Um Evento de Domínio (Domain Event) é um registro imutável que captura e expressa a ocorrência de um acontecimento relevante e com significado de negócio que já ocorreu no passado dentro do domínio.",
        ("Características essenciais: ", ""),
        ("• Verbo no particípio passado: ", "Expressa um fato consumado e inegável (ex.: PedidoCriado, PedidoPago, ClienteCadastrado)."),
        ("• Imutabilidade: ", "Fatos históricos não podem ser alterados ou deletados retroativamente; qualquer correção futura exige um novo evento de compensação."),
        ("• Carga útil (Payload): ", "Transporta os dados contextuais da ocorrência (identificador do agregado, timestamp occurredOn e atributos da alteração)."),
        ("• Desacoplamento e Consistência Eventual: ", "Permite notificar outros agregados ou microsserviços de forma assíncrona, eliminando acoplamentos temporais e chamadas bloqueantes.")
    ]
)

# QUESTÃO 9
code_q9 = """package br.edu.infnet.domain.event;

import java.time.Instant;
import java.util.UUID;

/**
 * Abstração que define o contrato base para todos os eventos de domínio da aplicação.
 */
public interface DomainEvent {

    /**
     * Identificador exclusivo da ocorrência do evento (rastreabilidade e idempotência).
     */
    UUID getEventId();

    /**
     * Carimbo de data/hora (UTC) do exato instante em que o fato de negócio ocorreu.
     */
    Instant getOccurredOn();
}"""

add_question(
    doc, 9,
    "Crie um trecho de código Java que mostre uma abstração de um objeto do tipo “evento de domínio”.",
    [
        "Abaixo está a interface DomainEvent, que padroniza o contrato fundamental de todos os eventos da aplicação com identificador único e timestamp:"
    ],
    code_snippets=[code_q9],
    note="https://github.com/Marcus-Boni/INFNET-26E2-26E3/blob/main/26E3_4/TP1/src/main/java/br/edu/infnet/domain/event/DomainEvent.java"
)

# QUESTÃO 10
code_q10 = """package br.edu.infnet.domain.event;

import java.math.BigDecimal;
import java.time.Instant;
import java.util.UUID;

public record PedidoPagoEvent(
    UUID eventId,
    Instant occurredOn,
    UUID pedidoId,
    UUID clienteId,
    BigDecimal valorPago
) implements DomainEvent {

    /**
     * Construtor de conveniência que gera ID e data/hora automaticamente na criação.
     */
    public PedidoPagoEvent(UUID pedidoId, UUID clienteId, BigDecimal valorPago) {
        this(UUID.randomUUID(), Instant.now(), pedidoId, clienteId, valorPago);
    }

    @Override
    public UUID getEventId() {
        return eventId;
    }

    @Override
    public Instant getOccurredOn() {
        return occurredOn;
    }
}"""

add_question(
    doc, 10,
    "Crie um trecho de código Java que mostre a implementação de um “evento de domínio” dentro do escopo do projeto Pet Friends.",
    [
        "Abaixo está a implementação concreta do evento PedidoPagoEvent, utilizando Java Record para assegurar imutabilidade nativa e semântica expressiva:"
    ],
    code_snippets=[code_q10],
    note="https://github.com/Marcus-Boni/INFNET-26E2-26E3/blob/main/26E3_4/TP1/src/main/java/br/edu/infnet/domain/event/PedidoPagoEvent.java"
)

# QUESTÃO 11
headers_q11 = ["Critério", "Fila (Queue)", "Tópico (Topic)"]
rows_q11 = [
    ["Padrão de Distribuição", "Ponto a Ponto (1 para 1)", "Publicação / Assinatura (1 para N)"],
    ["Destino da Mensagem", "Entregue a um único trabalhador (processou, remove)", "Entregue a todos os assinantes ativos (cópia para cada)"],
    ["Finalidade Principal", "Balanceamento de tarefas e carga concorrente", "Notificação de eventos para múltiplos serviços interessados"]
]

add_question(
    doc, 11,
    "Qual é a diferença entre filas e tópicos e como estes elementos funcionam em conjunto?",
    [
        "Diferença Fundamental:",
        ("• Fila (Queue - Ponto a Ponto): ", "Uma mensagem enviada para uma fila é entregue a apenas um consumidor entre os trabalhadores conectados. Quando esse consumidor conclui e confirma o processamento (ACK), a mensagem é removida. Ideal para distribuição de trabalho concorrente."),
        ("• Tópico (Topic - Publicação/Assinatura): ", "Uma mensagem enviada a um tópico é copiada e enviada a todos os assinantes inscritos (fan-out / broadcast). Ideal para notificação de fatos que interessam a vários subsistemas."),
        "Como funcionam em conjunto:",
        "Eles operam de maneira combinada no padrão Fan-out com Filas Dedicadas (como em RabbitMQ via Exchanges ou AWS SNS + SQS):",
        ("1. ", "O produtor emite o evento de domínio para um Tópico central."),
        ("2. ", "O Tópico replica esse evento e o encaminha para as Filas exclusivas de cada serviço consumidor interessado (ex.: Fila de Faturamento, Fila de Notificações)."),
        ("3. ", "Cada serviço consome de sua própria fila no seu próprio ritmo. Se o serviço de Notificações cair, sua fila retém as mensagens sem afetar o Faturamento ou os demais consumidores.")
    ],
    table_data=(headers_q11, rows_q11)
)

# QUESTÃO 12
diagram_q12 = """+----------------------------------------------------------------------------------------+
|                               MICROSSERVIÇO DE PEDIDOS                                 |
|                                                                                        |
|   [Requisição HTTP]                                                                    |
|           │                                                                            |
|           ▼                                                                            |
|   [PedidoService] ──► [Agregado Pedido]                                                |
|           │                 │                                                          |
|           │                 ▼                                                          |
|           │          Gera PedidoPagoEvent                                              |
|           │                                                                            |
|           ▼ (Mesma Transação ACID do Banco de Dados Relacional)                        |
|   ┌────────────────────────────────────────────────────────┐                           |
|   │                  BANCO DE DADOS LOCAL                  │                           |
|   │  ┌─────────────────────────┐ ┌──────────────────────┐  │                           |
|   │  │    Tabela tb_pedido     │ │  Tabela tb_outbox    │  │                           |
|   │  │  (Estado Atual Pedido)  │ │ (Eventos pendentes)  │  │                           |
|   │  └─────────────────────────┘ └──────────┬───────────┘  │                           |
|   └─────────────────────────────────────────┼──────────────┘                           |
|                                             │                                          |
|                                     Lê registros pendentes                             |
|                                             ▼                                          |
|                             [Outbox Worker / Debezium CDC]                             |
+─────────────────────────────────────────────┼──────────────────────────────────────────+
                                              │ Publica Mensagem
                                              ▼
                             ┌──────────────────────────────────┐
                             │       TÓPICO: pedidos-eventos    │ (Message Broker:
                             │      [PedidoPagoEvent Payload]   │  Kafka / RabbitMQ)
                             └────────────────┬─────────────────┘
                                              │
                      ┌───────────────────────┴───────────────────────┐
                      ▼                                               ▼
     ┌─────────────────────────────────┐             ┌─────────────────────────────────┐
     │      FILA: faturamento-pedidos  │             │     FILA: notificacao-pedidos   │
     └────────────────┬────────────────┘             └────────────────┬────────────────┘
                      ▼                                               ▼
         [Serviço de Faturamento]                        [Serviço de Notificação]
       (Emite Nota Fiscal e Cobrança)                 (Envia Confirmação ao Cliente)"""

add_question(
    doc, 12,
    "Dê um exemplo de solução de arquitetura para publicação de eventos de domínio dentro do escopo do projeto Pet Friends (ideal um desenho).",
    [
        "A solução de referência para garantir a publicação confiável e sem perda de dados em arquiteturas distribuídas é o padrão Transactional Outbox associado a um Message Broker (como Apache Kafka ou RabbitMQ):",
        "Desenho da Solução Arquitetural:"
    ],
    code_snippets=[diagram_q12],
    note="Funcionamento: 1) O PedidoService grava a atualização na tabela tb_pedido e o evento na tabela tb_outbox sob a mesma transação ACID local; 2) O Outbox Worker/CDC lê a tabela e publica no Tópico do Message Broker; 3) O Broker encaminha para as filas exclusivas dos serviços de Faturamento e Notificação."
)

# QUESTÃO 13
add_question(
    doc, 13,
    "Explique qual a finalidade de uma Event Store no contexto de eventos de domínio.",
    [
        "Uma Event Store é um banco de dados especializado projetado exclusivamente para armazenar, versionar e consultar sequências imutáveis de eventos de domínio ordenadas cronologicamente (event streams).",
        ("Suas finalidades primordiais são: ", ""),
        ("1. Fonte Única da Verdade (Single Source of Truth): ", "Ela não armazena tabelas mutáveis com o estado atual atualizado por UPDATE. O histórico completo e imutável de eventos gravados no stream do agregado é a verdade definitiva do negócio."),
        ("2. Persistência Estritamente Incremental (Append-Only): ", "Garante que eventos nunca sofram modificação ou deleção. Apenas inserções são permitidas, fornecendo trilha de auditoria contábil completa e nativa."),
        ("3. Controle de Concorrência Otimista por Stream: ", "Cada evento possui um número sequencial de versão. Se duas operações tentarem gravar concorrentemente o evento de versão N, a Event Store detecta o conflito e aborta a segunda transação, preservando a coerência."),
        ("4. Alimentação de Modelos de Leitura (CQRS): ", "Fornece assinaturas em tempo real (stream subscriptions) que alimentam bancos de dados otimizados para leitura e dashboards analíticos.")
    ]
)

# QUESTÃO 14
add_question(
    doc, 14,
    "O que é Event Sourcing e como ele se diferencia da persistência tradicional em bancos de dados relacionais? Explique como os eventos salvos são usados para recuperar o estado atual de um Agregado.",
    [
        ("1. O que é Event Sourcing e Diferença da Persistência Tradicional: ", ""),
        ("• Persistência Tradicional (CRUD / Orientada a Estado): ", "Armazena apenas o estado atual da entidade em colunas de uma tabela. A cada alteração, executa-se um UPDATE, sobrescrevendo o registro anterior e perdendo todo o histórico de transições e as intenções de negócio."),
        ("• Event Sourcing (Orientada a Eventos): ", "Em vez de persistir o estado final, armazena uma série cronológica de eventos de domínio imutáveis que documentam cada fato ocorrido desde o início. O banco de dados nunca sofre atualizações destrutivas."),
        ("2. Como os eventos salvos recuperam o estado atual do Agregado: ", ""),
        ("Para restaurar o Agregado em memória, utiliza-se o Replay de Eventos (Hydration): ", "1) O repositório recupera na Event Store todos os eventos do agregado ordenados por versão; 2) Instancia-se a entidade em seu estado inicial; 3) Itera-se sobre os eventos invocando métodos de mutação interna (ex.: apply(event)), reconstruindo passo a passo os atributos do objeto; 4) Ao finalizar o último evento, o agregado está no estado atual mais recente pronto para novas operações."),
        ("3. Otimização com Snapshots: ", "Para evitar reprocessar milhares de eventos a cada leitura, gera-se periodicamente uma 'foto' estática do agregado (Snapshot a cada N eventos). Na recuperação, carrega-se o último Snapshot e aplicam-se apenas os eventos que ocorreram depois dele.")
    ]
)

doc.save("TP1/TP1_Respostas_DDD_Arquitetura.docx")
print("Arquivo docx criado com sucesso em: TP1/TP1_Respostas_DDD_Arquitetura.docx")
